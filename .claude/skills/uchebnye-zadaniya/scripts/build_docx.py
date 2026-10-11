#!/usr/bin/env python3
"""Собирает .docx из упрощённого Markdown по требованиям оформления Дианы.

Использование:
    python3 build_docx.py input.md output.docx [--title-page title.json]
"""
import argparse
import datetime
import json
import re
import sys

try:
    from docx import Document
    from docx.enum.section import WD_SECTION
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt
except ImportError:
    sys.exit("Нужен python-docx: pip install python-docx")

FONT = "Times New Roman"
SIZE = 14
TABLE_SIZE = 12
MARGIN = Cm(1.27)  # «узкие» поля Word

STUDENT = {
    "university": "Учреждение образования\n«Белорусский государственный педагогический\nуниверситет имени Максима Танка»",
    "faculty": "Институт инклюзивного образования",
    "student": "Матвеева Диана Алексеевна",
    "course_group": "студентка 3 курса, группы 280524",
    "speciality": "профиль: дошкольная логопедия",
    "city": "Минск",
}

INLINE = re.compile(r"(\*\*.+?\*\*|\*.+?\*)")


def set_font(run, size=SIZE, bold=None, italic=None):
    run.font.name = FONT
    run.font.size = Pt(size)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), FONT)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def fmt_paragraph(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    pf = p.paragraph_format
    pf.alignment = align
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)


def add_inline(p, text, size=SIZE, bold=False, italic=False):
    for part in INLINE.split(text):
        if not part:
            continue
        b, i = bold, italic
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            part, b = part[2:-2], True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            part, i = part[1:-1], True
        set_font(p.add_run(part), size, b, i)


def para(doc, text="", align=WD_ALIGN_PARAGRAPH.JUSTIFY, bold=False, italic=False, size=SIZE):
    p = doc.add_paragraph()
    fmt_paragraph(p, align)
    add_inline(p, text, size, bold, italic)
    return p


def setup_document():
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = FONT
    st.font.size = Pt(SIZE)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    st.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    st.paragraph_format.space_after = Pt(0)
    st.paragraph_format.space_before = Pt(0)
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    for side in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(sec, side, MARGIN)
    return doc


def add_page_number(section):
    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    fmt_paragraph(p, WD_ALIGN_PARAGRAPH.CENTER)
    run = p.add_run()
    set_font(run)
    for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tag)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)


def title_page(doc, info):
    c = WD_ALIGN_PARAGRAPH.CENTER
    for line in STUDENT["university"].split("\n"):
        para(doc, line, c)
    para(doc, "")
    para(doc, STUDENT["faculty"], c)
    if info.get("kafedra"):
        para(doc, info["kafedra"], c)
    for _ in range(11):
        para(doc, "")
    if info.get("work_type"):
        para(doc, info["work_type"], c, bold=True)
    if info.get("discipline"):
        para(doc, info["discipline"], c)
    if info.get("topic"):
        para(doc, info["topic"], c, bold=True)
    for _ in range(8):
        para(doc, "")
    r = WD_ALIGN_PARAGRAPH.RIGHT
    para(doc, "Выполнила:", r)
    para(doc, STUDENT["course_group"], r)
    para(doc, STUDENT["speciality"], r)
    para(doc, STUDENT["student"], r)
    if info.get("teacher"):
        para(doc, "")
        para(doc, info["teacher"], r)
    for _ in range(10 if not info.get("teacher") else 8):
        para(doc, "")
    para(doc, f"{STUDENT['city']} {datetime.date.today().year}", c)
    # новая секция: на титульном листе нет номера страницы
    doc.add_section(WD_SECTION.NEW_PAGE)


def set_cell_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "000000")
        borders.append(el)
    tbl_pr.append(borders)


def add_table(doc, rows):
    cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
    cells = [r for r in cells if not all(re.fullmatch(r":?-{2,}:?", c) for c in r)]
    ncols = max(len(r) for r in cells)
    table = doc.add_table(rows=len(cells), cols=ncols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_cell_borders(table)
    for i, row in enumerate(cells):
        for j in range(ncols):
            cell = table.cell(i, j)
            p = cell.paragraphs[0]
            fmt_paragraph(p, WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT)
            add_inline(p, row[j] if j < len(row) else "", TABLE_SIZE, bold=(i == 0))
    para(doc, "")


def add_list_item(doc, marker, text):
    p = doc.add_paragraph()
    fmt_paragraph(p)
    p.paragraph_format.left_indent = Cm(0.63)
    p.paragraph_format.first_line_indent = Cm(-0.63)
    add_inline(p, f"{marker}\t{text}")
    p.paragraph_format.tab_stops.add_tab_stop(Cm(0.63))


def convert(md, doc):
    lines = md.splitlines()
    i, buf = 0, []

    def flush():
        if buf:
            para(doc, " ".join(s.strip() for s in buf))
            buf.clear()

    while i < len(lines):
        line = lines[i]
        s = line.strip()
        if not s:
            flush()
        elif s == r"\pagebreak":
            flush()
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        elif s.startswith("|"):
            flush()
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(lines[i])
                i += 1
            add_table(doc, rows)
            continue
        elif s.startswith("### "):
            flush()
            para(doc, s[4:], WD_ALIGN_PARAGRAPH.LEFT, bold=True, italic=True)
        elif s.startswith("## "):
            flush()
            text = s[3:]
            p = para(doc, text, WD_ALIGN_PARAGRAPH.CENTER if "источник" in text.lower() or "литератур" in text.lower() else WD_ALIGN_PARAGRAPH.JUSTIFY, bold=True)
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.keep_with_next = True
        elif s.startswith("# "):
            flush()
            p = para(doc, s[2:], WD_ALIGN_PARAGRAPH.CENTER, bold=True)
            p.paragraph_format.space_after = Pt(6)
        elif re.match(r"^[-*•]\s+", s):
            flush()
            add_list_item(doc, "–", re.sub(r"^[-*•]\s+", "", s))
        elif re.match(r"^\d+[.)]\s+", s):
            flush()
            m = re.match(r"^(\d+)[.)]\s+(.*)$", s)
            add_list_item(doc, f"{m.group(1)}.", m.group(2))
        else:
            buf.append(s)
        i += 1
    flush()


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("input")
    ap.add_argument("output")
    ap.add_argument("--title-page", help="JSON с полями титульного листа")
    a = ap.parse_args()

    doc = setup_document()
    if a.title_page:
        with open(a.title_page, encoding="utf-8") as f:
            title_page(doc, json.load(f))
        sec = doc.sections[-1]
        for side in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
            setattr(sec, side, MARGIN)
    add_page_number(doc.sections[-1])

    with open(a.input, encoding="utf-8") as f:
        convert(f.read(), doc)
    doc.save(a.output)
    print(f"Готово: {a.output}")


if __name__ == "__main__":
    main()
